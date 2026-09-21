#!/usr/bin/env python3
"""Workspace OS office-document generator.

Reads a JSON spec and writes a real .pptx / .xlsx / .docx using the
industry-standard python-pptx / openpyxl / python-docx libraries.

Usage:  gen.py <pptx|xlsx|docx> --spec <spec.json> --out <output-file>

On success the LAST stdout line is exactly:  WROTE:<absolute output path>
On failure a line beginning with ERROR: is printed to stderr and exit != 0.
"""
import argparse
import json
import os
import sys


def gen_pptx(spec, out):
    """Render a modern, themed deck from a spec.

    Styling lives in the pptx_design.Design "design system" (palette, type
    scale, layouts). This function owns spec → slide dispatch, back-compat
    layout aliases, live-metric anchors (`metrics`), slide numbers and notes.
    """
    from pptx import Presentation
    from pptx.util import Inches

    # Import the sibling design module whether run from its dir or via the shim.
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from pptx_design import Design  # noqa: E402

    design = Design(spec.get("theme", {}))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # Back-compat aliases: old specs said "title"; new specs say "cover".
    ALIAS = {"title": "cover", "bullet": "bullets"}
    # Full-bleed layouts own the whole slide (no page-number chrome).
    FULL_BLEED = {"cover", "section", "closing"}

    n_content = 0
    for s in spec.get("slides", []):
        slide = prs.slides.add_slide(blank)
        raw = s.get("layout", "bullets")
        layout = ALIAS.get(raw, raw)

        if layout == "cover":
            design.cover(slide, s)
        elif layout == "section":
            design.section(slide, s)
        elif layout == "closing":
            design.closing(slide, s)
        elif layout == "table":
            design.table(slide, s)
        elif layout == "chart":
            design.chart(slide, s)
        elif layout == "image":
            if s.get("title"):
                design.slide_title(slide, s["title"])
            if s.get("imagePath") and os.path.exists(s["imagePath"]):
                # Centered, aspect-preserving (width only) — good for screenshots.
                pic = slide.shapes.add_picture(
                    s["imagePath"], Inches(2.67), Inches(1.95), width=Inches(8.0))
                # Optional clickable source caption under the shot.
                src = s.get("sourceUrl") or s.get("href")
                if src:
                    from pptx_design import TYPE
                    y = pic.top + pic.height + Inches(0.12)
                    design.text(slide, Inches(2.67), y, Inches(8.0), Inches(0.4),
                                s.get("sourceLabel", str(src)), TYPE["caption"],
                                design.pal["accent"], link=str(src))
        else:  # "bullets" (default)
            design.bullets(slide, s)

        # Live-metric anchors: each becomes a named text shape
        # (`wos-metric-<tag>`) that wos-metric can later locate + rewrite.
        _place_metrics(design, slide, s.get("metrics", []))

        if layout not in FULL_BLEED:
            n_content += 1
            design.page_number(slide, n_content)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])
    prs.save(out)


def _place_metrics(design, slide, metrics):
    """Places big "stat-card" text shapes, each a named live-metric anchor.

    A metric entry: {"tag":"revenue", "text":"1200", "label":"Revenue",
    "x":<in>, "y":<in>}. `tag` → shape name `wos-metric-<tag>` (the anchor
    wos-metric rewrites); `text` is the single literal run (fidelity floor).
    """
    from pptx.util import Inches
    from pptx.enum.text import PP_ALIGN
    from pptx_design import TYPE
    for i, m in enumerate(metrics):
        tag = str(m.get("tag", "")).strip()
        if not tag:
            continue
        x = Inches(m.get("x", 0.9 + (i % 3) * 4.2))
        y = Inches(m.get("y", 3.4))
        # The value: ONE run, shape named for the writer to find + rewrite.
        design.text(slide, x, y, Inches(3.8), Inches(1.0),
                    m.get("text", ""), TYPE["cover_title"],
                    design.pal["primary"], bold=True, align=PP_ALIGN.LEFT,
                    name=f"wos-metric-{tag}")
        if m.get("label"):
            design.text(slide, x, y + Inches(1.05), Inches(3.8), Inches(0.5),
                        str(m["label"]).upper(), TYPE["caption"],
                        design.pal["muted"], bold=True, align=PP_ALIGN.LEFT)


def gen_xlsx(spec, out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(); first = True
    for sh in spec.get("sheets", []):
        ws = wb.active if first else wb.create_sheet()
        ws.title = sh.get("name", "Sheet"); first = False
        if sh.get("headers"):
            ws.append(sh["headers"])
            fill = PatternFill("solid", fgColor="1F4E79")
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = fill

        def add(row, bold=False):
            ws.append(row)
            r = ws.max_row
            for ci in sh.get("currencyCols", []):
                ws.cell(r, ci).number_format = '#,##0 €'
            for ci in sh.get("percentCols", []):
                ws.cell(r, ci).number_format = '0.0%'
            # A "links" column holds URLs — make each a clickable hyperlink. In
            # the app the LOK hyperlink callback opens it in the Workspace browser.
            for ci in sh.get("linkCols", []):
                cell = ws.cell(r, ci)
                url = cell.value
                if isinstance(url, str) and url.strip():
                    cell.hyperlink = url.strip()
                    cell.style = "Hyperlink"
            if bold:
                for c in ws[r]:
                    c.font = Font(bold=True)

        for row in sh.get("rows", []):
            add(row)
        if sh.get("totalRow"):
            add(sh["totalRow"], bold=True)
        for col in range(1, len(sh.get("headers", [1])) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16
        if sh.get("freezeHeader"):
            ws.freeze_panes = "A2"
    wb.save(out)


def gen_docx(spec, out):
    from docx import Document
    from docx.enum.text import WD_BREAK

    d = Document()
    if spec.get("title"):
        d.add_heading(spec["title"], 0)
    for sec in spec.get("sections", []):
        if sec.get("pageBreak"):
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if sec.get("heading"):
            d.add_heading(sec["heading"], int(sec.get("level", 1)))
        for para in sec.get("paragraphs", []):
            d.add_paragraph(str(para))
        for b in sec.get("bullets", []):
            d.add_paragraph(str(b), style="List Bullet")
        if sec.get("table"):
            hdr = sec["table"]["headers"]; rows = sec["table"]["rows"]
            tbl = d.add_table(rows=1, cols=len(hdr))
            try:
                tbl.style = "Light Grid Accent 1"
            except Exception:
                pass
            for c, h in enumerate(hdr):
                tbl.rows[0].cells[c].paragraphs[0].add_run(str(h)).bold = True
            for row in rows:
                cells = tbl.add_row().cells
                for c, val in enumerate(row):
                    cells[c].text = str(val)
    d.save(out)


GENERATORS = {"pptx": gen_pptx, "xlsx": gen_xlsx, "docx": gen_docx}


def main():
    ap = argparse.ArgumentParser(prog="gen.py")
    ap.add_argument("type", choices=GENERATORS.keys())
    ap.add_argument("--spec", required=True, help="path to the JSON spec ('-' for stdin)")
    ap.add_argument("--out", required=True, help="output file path")
    args = ap.parse_args()
    try:
        raw = sys.stdin.read() if args.spec == "-" else open(args.spec, encoding="utf-8").read()
        spec = json.loads(raw)
        out = os.path.abspath(args.out)
        GENERATORS[args.type](spec, out)
        print("WROTE:" + out)
    except Exception as e:  # noqa: BLE001 — surface any failure clearly to the agent
        print("ERROR: " + str(e), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
